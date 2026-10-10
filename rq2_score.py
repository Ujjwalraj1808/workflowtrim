"""RQ2 scores: inter-annotator agreement, then precision / recall / F1 of the detector.

Usage: python rq2_score.py
Reads rq2/annotator_A.xlsx, rq2/annotator_B.xlsx (filled in), rq2/tool_recall.csv and, once the
annotators have discussed their disagreements, rq2/resolved.csv (columns: id, smell, label).
  1. Every label must be filled in (Y/N on "recall", TP/FP on "precision").
  2. Cohen's kappa between A and B, per smell and pooled, on both sheets.
  3. Disagreements without a valid label in resolved.csv are written to rq2/disagreements.csv
     and the script stops.
  4. Final label = the agreed label, or the resolved one. Then per smell:
     precision = TP / labelled findings ("precision" sheet)
     recall    = files the detector flags / files the annotators labelled Y ("recall" sheet, file level)
"""
import csv
from collections import Counter
import os
from openpyxl import load_workbook
from summarize import SMELLS

OUT = "rq2"


def read(who):
    wb = load_workbook(os.path.join(OUT, f"annotator_{who}.xlsx"), read_only=True)
    labels, missing = {}, []
    rows = list(wb["recall"].iter_rows(values_only=True))
    header = rows[0]
    for row in rows[1:]:
        rec = dict(zip(header, row))
        for s in SMELLS:
            v = str(rec[s] or "").strip().upper()
            if v in ("Y", "N"):
                labels[(rec["id"], s)] = v
            else:
                missing.append(f"{rec['id']} {s}")
    rows = list(wb["precision"].iter_rows(values_only=True))
    header = rows[0]
    for row in rows[1:]:
        rec = dict(zip(header, row))
        v = str(rec["label"] or "").strip().upper()
        if v in ("TP", "FP"):
            labels[(rec["id"], rec["smell"])] = v
        else:
            missing.append(rec["id"])
    if missing:
        raise SystemExit(f"annotator_{who}.xlsx: {len(missing)} labels missing or not Y/N/TP/FP, "
                         f"e.g. {', '.join(missing[:10])}")
    return labels


def kappa(pairs):
    """Cohen's kappa for a list of (label_a, label_b). None if both annotators used one label only."""
    n = len(pairs)
    if n == 0:
        return None
    cats = {x for p in pairs for x in p}
    po = sum(a == b for a, b in pairs) / n
    pe = sum((sum(a == c for a, _ in pairs) / n) * (sum(b == c for _, b in pairs) / n) for c in cats)
    return None if pe == 1 else (po - pe) / (1 - pe)


def fmt(x):
    return "n/a" if x is None else f"{x:.2f}"


def sheet(key):
    return "recall" if key[0].startswith("R") else "precision"


def valid(key, label):
    return label in (("Y", "N") if sheet(key) == "recall" else ("TP", "FP"))


def main():
    a, b = read("A"), read("B")
    if a.keys() != b.keys():
        raise SystemExit("annotator_A.xlsx and annotator_B.xlsx do not contain the same items.")
    resolved = {}
    path = os.path.join(OUT, "resolved.csv")
    if os.path.exists(path):
        with open(path, newline="", encoding="utf-8-sig") as f:
            resolved = {(r["id"], r["smell"]): (r["label"] or "").strip().upper() for r in csv.DictReader(f)}

    print("Inter-annotator agreement: Cohen's kappa and % agreement (kappa is low when one label dominates,")
    print("so report both)")
    print(f"  {'Smell':<8}{'recall: kappa':>15}{'agree':>8}{'precision: kappa':>19}{'agree':>8}")
    for s in [*SMELLS, "pooled"]:
        cells = []
        for sh in ("recall", "precision"):
            pairs = [(a[key], b[key]) for key in a if sheet(key) == sh and (s == "pooled" or key[1] == s)]
            agree = f"{sum(x == y for x, y in pairs) * 100 / len(pairs):.0f}%" if pairs else "n/a"
            cells += [fmt(kappa(pairs)), agree]
        print(f"  {s:<8}{cells[0]:>15}{cells[1]:>8}{cells[2]:>19}{cells[3]:>8}")

    disagreements = [key for key in a if a[key] != b[key]]
    open_items = [key for key in disagreements if not valid(key, resolved.get(key, ""))]
    if open_items:
        with open(os.path.join(OUT, "disagreements.csv"), "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["id", "smell", "label", "label_A", "label_B"])
            for key in sorted(open_items):
                w.writerow([key[0], key[1], resolved.get(key, ""), a[key], b[key]])
        raise SystemExit(f"\n{len(open_items)} of {len(disagreements)} disagreements have no valid label in "
                         f"{OUT}/resolved.csv (Y/N for R-ids, TP/FP for P-ids) -> {OUT}/disagreements.csv. "
                         f"Discuss them, fill in `label`, add them to {OUT}/resolved.csv and run again.")
    final = {key: (a[key] if a[key] == b[key] else resolved[key]) for key in a}
    n_res = Counter(sheet(k) for k in disagreements)
    print(f"\n  Resolved by discussion: {n_res['recall']} recall and {n_res['precision']} precision labels.")

    with open(os.path.join(OUT, "tool_recall.csv"), newline="", encoding="utf-8") as f:
        tool = {(r["id"], s): r[s] for r in csv.DictReader(f) for s in SMELLS}

    print("\nDetector accuracy")
    print(f"  {'Smell':<8}{'TP/labelled':>12}{'Precision':>11}{'found/actual':>14}{'Recall':>9}{'F1':>7}")
    tot = [0, 0, 0, 0]
    for s in SMELLS:
        tp = sum(1 for k, v in final.items() if sheet(k) == "precision" and k[1] == s and v == "TP")
        n = sum(1 for k in final if sheet(k) == "precision" and k[1] == s)
        actual = [k for k, v in final.items() if sheet(k) == "recall" and k[1] == s and v == "Y"]
        found = sum(tool[k] == "Y" for k in actual)
        p = tp / n if n else None
        r = found / len(actual) if actual else None
        f1 = 2 * p * r / (p + r) if p is not None and r is not None and p + r else None
        tot = [tot[0] + tp, tot[1] + n, tot[2] + found, tot[3] + len(actual)]
        print(f"  {s:<8}{f'{tp}/{n}':>12}{fmt(p):>11}{f'{found}/{len(actual)}':>14}{fmt(r):>9}{fmt(f1):>7}")
    p, r = tot[0] / tot[1], (tot[2] / tot[3] if tot[3] else None)
    f1 = 2 * p * r / (p + r) if r is not None and p + r else None
    print(f"  {'all':<8}{f'{tot[0]}/{tot[1]}':>12}{fmt(p):>11}{f'{tot[2]}/{tot[3]}':>14}{fmt(r):>9}{fmt(f1):>7}")
    fp_tool = sum(1 for k, v in final.items() if sheet(k) == "recall" and v == "N" and tool[k] == "Y")
    print(f"\n  Recall sheet: the detector also flagged {fp_tool} file-smell pairs the annotators labelled N.")


if __name__ == "__main__":
    main()
