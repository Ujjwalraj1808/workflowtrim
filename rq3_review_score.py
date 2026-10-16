"""RQ3 manual review: agreement between the two reviewers and the final verdict per smell.

Inputs (in rq3/): review_A.xlsx, review_B.xlsx (each reviewer's own copy of review.xlsx, column `label`:
correct / incorrect / unsure) and disagreements.xlsx (column `resolved_label` for every row where A and B
differ, decided by the two reviewers together). Writes rq3/review_resolved.csv and prints the tables.

Run:  python rq3_review_score.py
"""
import csv
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent / "rq3"
LABELS = ("correct", "incorrect", "unsure")
SMELLS = ("S1", "S2", "S3", "S5", "S6", "S8")


def rows(path):
    ws = load_workbook(path, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    return {r[0]: dict(zip(header, r)) for r in it if r[0]}


def label(value):
    v = str(value or "").strip().lower()
    if v not in LABELS:
        raise SystemExit(f"invalid label: {value!r}")
    return v


def kappa(a, b):
    """Cohen's kappa; None when chance agreement is 1 (all labels identical)."""
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in LABELS) / (n * n)
    return po, (None if pe == 1 else (po - pe) / (1 - pe))


def wilson(k, n, z=1.96):
    p = k / n
    centre = (p + z * z / (2 * n)) / (1 + z * z / n)
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return centre - half, centre + half


def main():
    a, b = rows(HERE / "review_A.xlsx"), rows(HERE / "review_B.xlsx")
    resolved = rows(HERE / "disagreements.xlsx")
    if set(a) != set(b):
        raise SystemExit("review_A and review_B do not contain the same rows")
    final = {}
    for rid in sorted(a):
        la, lb = label(a[rid]["label"]), label(b[rid]["label"])
        if la == lb:
            final[rid] = la
        elif rid in resolved:
            final[rid] = label(resolved[rid]["resolved_label"])
        else:
            raise SystemExit(f"{rid}: reviewers disagree and no resolved_label is given")

    with open(HERE / "review_resolved.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "smell", "repo", "file", "label_A", "label_B", "label"])
        for rid in sorted(final):
            w.writerow([rid, a[rid]["smell"], a[rid]["repo"], a[rid]["file"],
                        label(a[rid]["label"]), label(b[rid]["label"]), final[rid]])

    print(f"{'smell':6}{'n':>4}{'correct':>9}{'incorrect':>11}{'unsure':>8}{'agree%':>8}{'kappa':>7}")
    for s in SMELLS + ("all",):
        ids = [r for r in sorted(final) if s == "all" or a[r]["smell"] == s]
        c = Counter(final[r] for r in ids)
        po, k = kappa([label(a[r]["label"]) for r in ids], [label(b[r]["label"]) for r in ids])
        ks = "  n/a" if k is None else f"{k:7.2f}"
        print(f"{s:6}{len(ids):>4}{c['correct']:>9}{c['incorrect']:>11}{c['unsure']:>8}{100 * po:>8.1f}{ks}")
    n, k = len(final), sum(v == "correct" for v in final.values())
    lo, hi = wilson(k, n)
    print(f"\ncorrect: {k}/{n} = {100 * k / n:.1f}% (95% Wilson CI {100 * lo:.1f}-{100 * hi:.1f}%)")


if __name__ == "__main__":
    main()
